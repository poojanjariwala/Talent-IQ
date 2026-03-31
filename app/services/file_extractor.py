import io
import re
from typing import Optional


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF bytes using pypdf."""
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        pages = []
        for page in reader.pages:
            text = page.extract_text() or ""
            pages.append(text)
        return "\n".join(pages).strip()
    except Exception as e:
        raise ValueError(f"PDF extraction failed: {e}")


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX bytes using python-docx."""
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(paragraphs).strip()
    except Exception as e:
        raise ValueError(f"DOCX extraction failed: {e}")


def extract_text_from_txt(file_bytes: bytes) -> str:
    """Decode plain text file."""
    try:
        return file_bytes.decode("utf-8", errors="replace").strip()
    except Exception as e:
        raise ValueError(f"TXT extraction failed: {e}")


def extract_text_from_csv(file_bytes: bytes) -> str:
    """Extract text from CSV (treat each row as text)."""
    import csv
    try:
        text = file_bytes.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows = []
        for row in reader:
            rows.append(" | ".join(row))
        return "\n".join(rows).strip()
    except Exception as e:
        raise ValueError(f"CSV extraction failed: {e}")


def extract_text_from_xlsx(file_bytes: bytes) -> str:
    """Extract text from XLSX spreadsheet."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        rows = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join(str(v) for v in row if v is not None)
                if row_text.strip():
                    rows.append(row_text)
        return "\n".join(rows).strip()
    except Exception as e:
        raise ValueError(f"XLSX extraction failed: {e}")


def extract_text(filename: str, file_bytes: bytes) -> str:
    """Dispatch extraction based on file extension."""
    ext = filename.lower().rsplit(".", 1)[-1]
    extractors = {
        "pdf": extract_text_from_pdf,
        "docx": extract_text_from_docx,
        "doc": extract_text_from_docx,
        "txt": extract_text_from_txt,
        "csv": extract_text_from_csv,
        "xlsx": extract_text_from_xlsx,
        "xls": extract_text_from_xlsx,
    }
    extractor = extractors.get(ext)
    if extractor is None:
        raise ValueError(f"Unsupported file type: .{ext}")
    return extractor(file_bytes)


def clean_text(text: str) -> str:
    """Remove excessive whitespace while preserving structure."""
    text = re.sub(r"\r\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r" {2,}", " ", text)
    return text.strip()
